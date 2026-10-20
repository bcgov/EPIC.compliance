"""Test ceb_summary report service."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from faker import Faker
from openpyxl import load_workbook
import pandas as pd
import pytest

from compliance_api.models import db
from compliance_api.models.agency import Agency
from compliance_api.models.case_file import CaseFile
from compliance_api.models.complaint.complaint import Complaint, ComplaintStatusEnum
from compliance_api.models.complaint.complaint_option import ComplaintSource, ComplaintSourceEnum
from compliance_api.models.complaint.complaint_resolution import ComplaintResolutionEnum
from compliance_api.models.inspection.inspection_req_enforcement_map import InspectionReqEnforcementMap
from compliance_api.models.inspection.inspection_attendance import InspectionAttendance
from compliance_api.models.inspection.inspection_firstnation import InspectionFirstnation
from compliance_api.models.inspection.inspection_option import InspectionAttendanceOption
from compliance_api.models.inspection_record import InspectionRecord
from compliance_api.services.report import ceb_summary
from compliance_api.models.inspection.inspection import Inspection
from compliance_api.models.inspection.inspection_requirement import InspectionRequirement
from compliance_api.models.topic import Topic
from compliance_api.models.compliance_finding import ComplianceFindingOption
from compliance_api.models.staff_user import StaffUser
from compliance_api.utils.datetime import BC_TIMEZONE

fake = Faker()


class TestCEBSummaryReportGenerator:
    """Test CEB Summary Report Generator."""

    @pytest.fixture(autouse=True)
    def setup(self, mocker):
        """Fixture to execute before and after each test."""
        mocker.patch(
            "compliance_api.services.report.ceb_summary.TrackService.get_first_nations",
            return_value=[],
        )
        self.insp_req = self._create_test_inspection_requirement()
        self.complaints = []
        self.agencies = []
        yield
        self._clean_up_database()

    def test_build_inspections_tab_query(self):
        """Test building inspections tab query with no date range."""
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=100),
        })
        results = generator._build_inspections_tab_query().all()
        assert len(results) == 1
        assert results[0].InspectionRequirement == self.insp_req

    def test_build_inspections_tab_query_results_expected_within_date_range(self):
        """Test building inspections tab query with date range that includes data."""
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=150),
            "end_date": datetime.now() + timedelta(days=153)
        })
        results = generator._build_inspections_tab_query().all()
        assert results[0].InspectionRequirement == self.insp_req

    def test_build_inspections_tab_query_no_results_expected_with_start_date(self):
        """Test building inspections tab query with start date that excludes data."""
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=154)
        })
        results = generator._build_inspections_tab_query().all()
        assert len(results) == 0

    def test_build_inspections_tab_query_no_results_expected_with_end_date(self):
        """Test building inspections tab query with end date that excludes data."""
        generator = ceb_summary.CEBSummaryReportGenerator({"end_date": datetime.now() - timedelta(days=10)})
        results = generator._build_inspections_tab_query().all()
        assert len(results) == 0

    def test_build_complaints_tab_query_includes_complaint_received_within_date_range(self):
        """Test complaints tab query includes complaints received within the date range."""
        complaint = self._create_test_complaint(datetime.now() + timedelta(days=152))
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=150),
            "end_date": datetime.now() + timedelta(days=153)
        })
        results = generator._build_complaints_tab_query().all()
        assert complaint.complaint_number in [row.complaint_number for row in results]

    def test_build_complaints_tab_query_excludes_complaint_received_before_start_date(self):
        """Test complaints tab query excludes complaints received before the start date."""
        complaint = self._create_test_complaint(datetime.now() + timedelta(days=149))
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=150),
            "end_date": datetime.now() + timedelta(days=153)
        })
        results = generator._build_complaints_tab_query().all()
        assert complaint.complaint_number not in [row.complaint_number for row in results]

    def test_build_complaints_tab_query_excludes_complaint_received_after_end_date(self):
        """Test complaints tab query excludes complaints received after the end date."""
        complaint = self._create_test_complaint(datetime.now() + timedelta(days=154))
        generator = ceb_summary.CEBSummaryReportGenerator({
            "start_date": datetime.now() + timedelta(days=150),
            "end_date": datetime.now() + timedelta(days=153)
        })
        results = generator._build_complaints_tab_query().all()
        assert complaint.complaint_number not in [row.complaint_number for row in results]

    def test_format_complaints_tab_data_includes_response_dates_turnaround_and_referred_agency(self):
        """Test response dates, calendar-day turnaround and referred agency are formatted for a complaint."""
        agency = self._create_test_agency()
        complaint = self._create_test_complaint(
            datetime(2026, 3, 1, 10, 0, tzinfo=BC_TIMEZONE),
            initial_response_date=datetime(2026, 3, 4, 9, 0, tzinfo=BC_TIMEZONE),
            final_response_date=datetime(2026, 3, 31, 16, 0, tzinfo=BC_TIMEZONE),
            resolution_id=ComplaintResolutionEnum.REFERRED_TO_AGENCY.value,
            resolution_agency_id=agency.id,
        )

        item = self._format_single_complaint(complaint)

        assert item["initial_response_date"] == "2026-03-04"
        assert item["final_response_date"] == "2026-03-31"
        assert item["days_to_initial_response"] == 3
        # Spans March DST change
        assert item["days_to_final_response"] == 30
        assert item["referred_agency"] == agency.name

    def test_format_complaints_tab_data_blank_when_response_dates_missing(self):
        """Test missing response dates leave dates and turnaround blank, and each interval is independent."""
        complaint = self._create_test_complaint(
            datetime(2026, 3, 1, 10, 0, tzinfo=BC_TIMEZONE),
            final_response_date=datetime(2026, 3, 11, 10, 0, tzinfo=BC_TIMEZONE),
        )

        item = self._format_single_complaint(complaint)

        assert item["initial_response_date"] is None
        assert item["days_to_initial_response"] is None
        assert item["days_to_final_response"] == 10
        assert item["referred_agency"] is None

    def test_format_complaints_tab_data_referred_agency_blank_for_other_resolutions(self):
        """Test referred agency is blank when the resolution is not Referred to Another Agency."""
        agency = self._create_test_agency()
        complaint = self._create_test_complaint(
            datetime(2026, 3, 1, 10, 0, tzinfo=BC_TIMEZONE),
            resolution_id=ComplaintResolutionEnum.UNFOUNDED.value,
            resolution_agency_id=agency.id,
        )

        item = self._format_single_complaint(complaint)

        assert item["referred_agency"] is None

    def test_format_complaints_tab_data_referred_agency_blank_when_no_agency(self):
        """Test referred agency is blank when the complaint was referred but no agency was specified."""
        complaint = self._create_test_complaint(
            datetime(2026, 3, 1, 10, 0, tzinfo=BC_TIMEZONE),
            resolution_id=ComplaintResolutionEnum.REFERRED_TO_AGENCY.value,
        )

        item = self._format_single_complaint(complaint)

        assert item["referred_agency"] is None

    def test_calendar_days_between_uses_bc_local_dates(self):
        """Test turnaround counts BC calendar days, not UTC days."""
        # UTC, as Postgres returns them: same UTC date, but 23:30 and 00:30 BC on consecutive BC dates
        received = datetime(2026, 10, 2, 6, 30, tzinfo=timezone.utc)
        responded = datetime(2026, 10, 2, 7, 30, tzinfo=timezone.utc)

        assert ceb_summary.CEBSummaryReportGenerator._calendar_days_between(received, responded) == 1
        assert ceb_summary.CEBSummaryReportGenerator._calendar_days_between(received, received) == 0
        assert ceb_summary.CEBSummaryReportGenerator._calendar_days_between(received, None) is None
        assert ceb_summary.CEBSummaryReportGenerator._calendar_days_between(None, responded) is None

    def test_to_excel_writes_complaint_columns_and_turnaround_section(self):
        """Test the workbook has the new Complaints columns and a turnaround table on Complaints Insights."""
        complaint = self._create_test_complaint(
            datetime(2026, 3, 1, 10, 0, tzinfo=BC_TIMEZONE),
            initial_response_date=datetime(2026, 3, 4, 9, 0, tzinfo=BC_TIMEZONE),
        )
        generator = ceb_summary.CEBSummaryReportGenerator({})
        rows = [r for r in generator._build_complaints_tab_query().all()
                if r.complaint_number == complaint.complaint_number]
        workbook = self._render_workbook(generator, generator._format_complaints_tab_data(rows))

        complaints_sheet = workbook["Complaints"]
        headers = [cell.value for cell in complaints_sheet[1]]
        assert headers[4:7] == ["Date Received", "Initial Response", "Final Response"]
        assert headers[11:13] == ["Complaint Resolution", "Referred Agency"]
        values = dict(zip(headers, [cell.value for cell in complaints_sheet[2]]))
        assert values["Initial Response"] == "2026-03-04"
        assert values["Final Response"] is None

        insights = workbook["ComplaintsInsights"]
        table = insights.tables["ComplaintResponseTurnaround"]
        first_cell = insights[table.ref.split(":")[0]]
        title = insights.cell(row=2, column=first_cell.column).value
        assert title == "Complaint Response Turnaround (Calendar Days)"
        header_row = [insights.cell(row=first_cell.row, column=first_cell.column + i).value for i in range(7)]
        data_row = [insights.cell(row=first_cell.row + 1, column=first_cell.column + i).value for i in range(7)]
        assert header_row[5:] == ["Received → Initial Response (Days)", "Received → Final Response (Days)"]
        assert data_row[0] == complaint.complaint_number
        assert data_row[5:] == [3, None]
        border = insights.cell(row=first_cell.row, column=first_cell.column).border
        assert (border.left.style, border.top.style, border.right.style, border.bottom.style) == (
            "thin", "thin", None, None
        )
        last_cell = insights[table.ref.split(":")[1]]
        border = last_cell.border
        assert (border.left.style, border.top.style, border.right.style, border.bottom.style) == (
            None, None, "thin", "thin"
        )

    def test_turnaround_section_does_not_overlap_pivots(self):
        """Test the turnaround table starts to the right of every Complaints Insights pivot."""
        generator = ceb_summary.CEBSummaryReportGenerator({})
        workbook = self._render_workbook(generator, [])

        insights = workbook["ComplaintsInsights"]
        table = insights.tables["ComplaintResponseTurnaround"]
        table_start_col = insights[table.ref.split(":")[0]].column
        for pivot in insights._pivots:
            pivot_end_col = insights[pivot.location.ref.split(":")[1]].column
            assert table_start_col > pivot_end_col

    def _format_single_complaint(self, complaint):
        """Run the complaints tab query and formatter, returning the item for the given complaint."""
        generator = ceb_summary.CEBSummaryReportGenerator({})
        results = [r for r in generator._build_complaints_tab_query().all()
                   if r.complaint_number == complaint.complaint_number]
        items = generator._format_complaints_tab_data(results)
        return next(item for item in items if item["complaint_number"] == complaint.complaint_number)

    @staticmethod
    def _render_workbook(generator, complaints_data):
        """Render the CEB workbook with only the given complaints data and load it back."""
        _, inspections_columns = generator._get_inspections_tab_columns_and_headers()
        _, enforcements_columns = generator._get_enforcements_tab_columns_and_headers()
        _, requirements_columns = generator._get_requirements_tab_columns_and_headers()
        complaints_headers, complaints_columns = generator._get_complaints_tab_columns_and_headers()
        inspections_headers, _ = generator._get_inspections_tab_columns_and_headers()
        enforcements_headers, _ = generator._get_enforcements_tab_columns_and_headers()
        requirements_headers, _ = generator._get_requirements_tab_columns_and_headers()
        output = generator._to_excel(
            pd.DataFrame(), inspections_columns, inspections_headers,
            pd.DataFrame(), enforcements_columns, enforcements_headers,
            pd.DataFrame(), requirements_columns, requirements_headers,
            pd.json_normalize(complaints_data), complaints_columns, complaints_headers,
        )
        return load_workbook(BytesIO(output))

    def _create_test_agency(self):
        """Create an agency for referral tests."""
        agency = Agency(name=fake.pystr(min_chars=5, max_chars=10))
        db.session.add(agency)
        db.session.flush()
        self.agencies.append(agency)
        return agency

    def _create_test_complaint(self, date_received, **extra_fields):
        """Create a complaint received on the given date."""
        case_file = CaseFile(
            date_created=datetime.now(),
            case_file_number=fake.pystr(min_chars=5, max_chars=10),
            initiation_id=1
        )
        db.session.add(case_file)
        db.session.flush()

        complaint_source = db.session.query(ComplaintSource).filter(
            ComplaintSource.name == ComplaintSourceEnum.PUBLIC.value
        ).first()
        if not complaint_source:
            complaint_source = ComplaintSource(name=ComplaintSourceEnum.PUBLIC.value)
            db.session.add(complaint_source)
            db.session.flush()

        complaint = Complaint(
            case_file_id=case_file.id,
            date_received=date_received,
            source_type_id=complaint_source.id,
            concern_description=fake.text(max_nb_chars=200),
            status=ComplaintStatusEnum.OPEN,
            complaint_number=fake.pystr(min_chars=5, max_chars=10),
            **extra_fields,
        )
        db.session.add(complaint)
        db.session.commit()
        self.complaints.append(complaint)

        return complaint

    def _create_test_inspection_requirement(self):
        """Create an inspection requirement for testing."""
        case_file = CaseFile(
            date_created=datetime.now(),
            case_file_number=fake.pystr(min_chars=5, max_chars=10),
            initiation_id=1
        )

        db.session.add(case_file)
        db.session.flush()

        topic = Topic(
            name="Test Topic"
        )
        finding = ComplianceFindingOption(
            name=fake.pystr(min_chars=5, max_chars=10)
        )
        officer = StaffUser(
            first_name=fake.pystr(min_chars=5, max_chars=10),
            last_name=fake.last_name(),
            position_id=1
        )
        inspection = Inspection(
            ir_number=fake.pystr(min_chars=5, max_chars=10),
            primary_officer=officer,
            start_date=datetime.now() + timedelta(days=152),
            end_date=datetime.now() + timedelta(days=152),
            initiation_id=1,
            case_file_id=case_file.id
        )

        db.session.add_all([topic, finding, officer, inspection, case_file])
        db.session.flush()

        inspection_record = InspectionRecord(
            inspection_id=inspection.id,
            date_issued=datetime.now(),
            ir_status_id=1
        )

        db.session.add(inspection_record)
        db.session.flush()

        # Get existing attendance option (seeded data) for INNER JOIN in query
        attendance_option = db.session.query(InspectionAttendanceOption).first()

        inspection_attendance = InspectionAttendance(
            inspection_id=inspection.id,
            attendance_option_id=attendance_option.id
        )
        inspection_firstnation = InspectionFirstnation(
            inspection_id=inspection.id,
            firstnation_id=1
        )
        db.session.add_all([inspection_attendance, inspection_firstnation])
        db.session.flush()

        insp_req = InspectionRequirement(
            inspection_id=inspection.id,
            topic_id=topic.id,
            compliance_finding_id=finding.id,
            summary="Requirement 1",
            sort_order=1,
        )

        db.session.add(insp_req)
        db.session.flush()

        insp_req_enf_map = InspectionReqEnforcementMap(
            requirement_id=insp_req.id,
            enforcement_action_id=1
        )

        db.session.add(insp_req_enf_map)
        db.session.commit()

        return insp_req

    def _clean_up_database(self):
        """Clean up the database after tests."""
        db.session.query(InspectionReqEnforcementMap).where(
            InspectionReqEnforcementMap.requirement_id == self.insp_req.id
        ).delete()
        db.session.query(InspectionRequirement).where(InspectionRequirement.id == self.insp_req.id).delete()
        db.session.query(Complaint).where(Complaint.id.in_([c.id for c in self.complaints])).delete()
        db.session.query(Agency).where(Agency.id.in_([a.id for a in self.agencies])).delete()
        db.session.commit()
